from __future__ import annotations

import csv
import html as html_lib
import io
import json
import math
import re
import time
import unicodedata
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd
import pdfplumber
from lxml import html
from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = ROOT / "data" / "raw"
PROCESSED_DIR = ROOT / "data" / "processed"
AUDIT_DIR = ROOT / "data" / "audit"
WIPO_FILE = RAW_DIR / "wipo-pct-2024.xlsx"
WIPO_FACT_SHEET_FILE = RAW_DIR / "wipo-country-profile-pct-2024.pdf"
UN_ENERGY_FILE = RAW_DIR / "un-energy-statistics-pocketbook-2025.pdf"
ACCESS_DATE = "2026-09-21"

WORLD_BANK_COUNTRIES_URL = "https://api.worldbank.org/v2/country?format=json&per_page=400"
WORLD_BANK_API = "https://api.worldbank.org/v2/country/all/indicator/{indicator}?format=json&per_page=30000&date={start}:{end}"
UN_MEMBER_STATES_URL = "https://www.un.org/about-us/member-states"
UN_OBSERVER_STATES_URL = "https://www.un.org/en/about-us/non-member-states"
NATURE_BASE = "https://www.nature.com/nature-index/research-leaders/2025/country/all/{scope}"
CLARIVATE_LIST_URL = "https://clarivate.com/highly-cited-researchers/"
CLARIVATE_AJAX_URL = "https://clarivate.com/wp-admin/admin-ajax.php"
WIPO_URL = "https://www.wipo.int/edocs/pubdocs/en/wipo-pub-901-2025-tech1.xlsx"
WIPO_FACT_SHEET_URL = "https://www.wipo.int/edocs/statistics-country-profile/en/_list/l5.pdf"
UN_ENERGY_PDF_URL = "https://desapublications.un.org/file/21030/download"

USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) CDI-Coverage-Audit/1.0"

WB_INDICATORS = {
    "gni": {"code": "NY.GNP.PCAP.PP.CD", "target": 2024, "start": 2019, "max_lag": 2},
    "life": {"code": "SP.DYN.LE00.IN", "target": 2024, "start": 2019, "max_lag": 2},
    "internet": {"code": "IT.NET.USER.ZS", "target": 2024, "start": 2019, "max_lag": 5},
    "water": {"code": "SH.H2O.BASW.ZS", "target": 2024, "start": 2019, "max_lag": 5},
    "electricity": {"code": "EG.USE.ELEC.KH.PC", "target": 2023, "start": 2018, "max_lag": 2},
    # The WDI API stores the 2023 LPI edition under the 2022 survey year.
    "lpi": {"code": "LP.LPI.OVRL.XQ", "target": 2022, "display_year": 2023, "start": 2014, "max_lag": 8},
}

NATURE_SCOPES = [
    "regions-Asia%20Pacific",
    "regions-Europe",
    "regions-South%20America",
    "regions-North%20America",
    "regions-Western%20Asia",
    "regions-Africa",
]


NAME_ALIASES = {
    "bahamas": "BHS",
    "bahamas the": "BHS",
    "bolivia plurinational state of": "BOL",
    "bolivia": "BOL",
    "brunei darussalam": "BRN",
    "brunei": "BRN",
    "cape verde": "CPV",
    "cabo verde": "CPV",
    "china": "CHN",
    "china mainland": "CHN",
    "congo": "COG",
    "congo republic of": "COG",
    "democratic republic of the congo": "COD",
    "congo dem rep": "COD",
    "cote d ivoire": "CIV",
    "ivory coast": "CIV",
    "czech republic": "CZE",
    "czechia": "CZE",
    "egypt": "EGY",
    "egypt arab rep": "EGY",
    "gambia": "GMB",
    "gambia the": "GMB",
    "hong kong sar": "HKG",
    "hong kong sar china": "HKG",
    "iran": "IRN",
    "iran islamic republic of": "IRN",
    "iran islamic rep": "IRN",
    "korea republic of": "KOR",
    "republic of korea": "KOR",
    "korea rep": "KOR",
    "south korea": "KOR",
    "north korea": "PRK",
    "democratic people s republic of korea": "PRK",
    "kyrgyz republic": "KGZ",
    "kyrgyzstan": "KGZ",
    "lao people s democratic republic": "LAO",
    "lao pdr": "LAO",
    "laos": "LAO",
    "macao sar china": "MAC",
    "macau sar": "MAC",
    "micronesia federated states of": "FSM",
    "micronesia fed sts": "FSM",
    "micronesia": "FSM",
    "moldova republic of": "MDA",
    "republic of moldova": "MDA",
    "moldova": "MDA",
    "north macedonia": "MKD",
    "nauru": "NRU",
    "russia": "RUS",
    "russian federation": "RUS",
    "slovak republic": "SVK",
    "slovakia": "SVK",
    "saint kitts and nevis": "KNA",
    "st kitts and nevis": "KNA",
    "saint lucia": "LCA",
    "st lucia": "LCA",
    "saint vincent and the grenadines": "VCT",
    "st vincent and the grenadines": "VCT",
    "sao tome and principe": "STP",
    "syrian arab republic": "SYR",
    "syria": "SYR",
    "taiwan": "TWN",
    "tanzania united republic of": "TZA",
    "united republic of tanzania": "TZA",
    "tanzania": "TZA",
    "turkey": "TUR",
    "turkiye": "TUR",
    "united kingdom uk": "GBR",
    "united kingdom": "GBR",
    "united kingdom of great britain and northern ireland": "GBR",
    "uk": "GBR",
    "united states of america usa": "USA",
    "united states of america": "USA",
    "united states": "USA",
    "us": "USA",
    "venezuela bolivarian republic of": "VEN",
    "venezuela rb": "VEN",
    "venezuela": "VEN",
    "viet nam": "VNM",
    "vietnam": "VNM",
    "west bank and gaza": "PSE",
    "palestinian territories": "PSE",
    "state of palestine": "PSE",
    "yemen": "YEM",
    "yemen rep": "YEM",
    "kosovo": "XKX",
    "somalia": "SOM",
    "vatican": "VAT",
    "holy see": "VAT",
    "china hong kong special administrative region": "HKG",
    "china macao special administrative region": "MAC",
    "anguilla": "AIA",
    "bonaire sint eustatius and saba": "BES",
    "cook islands": "COK",
    "falkland islands malvinas": "FLK",
    "guernsey": "GGY",
    "jersey": "JEY",
    "montserrat": "MSR",
    "niue": "NIU",
    "saint barthelemy": "BLM",
    "saint helena": "SHN",
    "saint martin french part": "MAF",
    "saint pierre and miquelon": "SPM",
    "united states virgin islands": "VIR",
    "wallis and futuna islands": "WLF",
}


@dataclass
class Observation:
    value: float | int | None
    year: int | None
    status: str
    source: str
    note: str = ""


def ensure_dirs() -> None:
    for path in (RAW_DIR, PROCESSED_DIR, AUDIT_DIR):
        path.mkdir(parents=True, exist_ok=True)


def fetch_bytes(url: str, retries: int = 3) -> bytes:
    last_error: Exception | None = None
    for attempt in range(retries):
        try:
            request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
            with urllib.request.urlopen(request, timeout=60) as response:
                return response.read()
        except Exception as error:  # noqa: BLE001
            last_error = error
            time.sleep(1.5 * (attempt + 1))
    raise RuntimeError(f"Failed to fetch {url}") from last_error


def fetch_text(url: str) -> str:
    return fetch_bytes(url).decode("utf-8", errors="replace")


def fetch_json(url: str) -> Any:
    return json.loads(fetch_bytes(url).decode("utf-8"))


def normalize_name(value: str) -> str:
    value = html_lib.unescape(value or "")
    value = re.sub(r"\s*\([a-d]\)\s*$", "", value, flags=re.IGNORECASE)
    value = re.sub(r"\s*\([^)]*\)\s*$", lambda m: " " + m.group(0).strip(" ()") if any(x in m.group(0) for x in ("USA", "UK")) else "", value)
    value = unicodedata.normalize("NFKD", value)
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    value = value.replace("&", " and ")
    value = re.sub(r"[^A-Za-z0-9]+", " ", value).strip().lower()
    return re.sub(r"\s+", " ", value)


def to_number(value: Any) -> float | int | None:
    if value is None:
        return None
    if isinstance(value, (int, float)) and not pd.isna(value):
        return int(value) if float(value).is_integer() else float(value)
    text = str(value).strip().replace(",", "")
    if not text or text.lower() in {"n.a.", "n.a", "na", "nan", ".."}:
        return None
    number = float(text)
    return int(number) if number.is_integer() else number


def build_master() -> list[dict[str, Any]]:
    payload = fetch_json(WORLD_BANK_COUNTRIES_URL)
    rows = []
    for item in payload[1]:
        region_id = item.get("region", {}).get("id")
        if not region_id or region_id == "NA":
            continue
        rows.append(
            {
                "iso3": item["id"],
                "iso2": item.get("iso2Code"),
                "name": item["name"],
                "region": item.get("region", {}).get("value"),
                "income": item.get("incomeLevel", {}).get("value"),
                "master_source": "World Bank country/economy list",
            }
        )

    extras = [
        {"iso3": "TWN", "iso2": "TW", "name": "Taiwan", "region": "East Asia & Pacific", "income": None, "master_source": "Additional ISO economy"},
        {"iso3": "VAT", "iso2": "VA", "name": "Vatican", "region": "Europe & Central Asia", "income": None, "master_source": "Additional ISO economy"},
    ]
    existing = {row["iso3"] for row in rows}
    rows.extend(row for row in extras if row["iso3"] not in existing)
    rows.sort(key=lambda row: row["iso3"])
    (RAW_DIR / "world-bank-countries.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return rows


def build_name_index(master: list[dict[str, Any]]) -> dict[str, str]:
    index = {normalize_name(row["name"]): row["iso3"] for row in master}
    index.update(NAME_ALIASES)
    return index


def resolve_name(name: str, name_index: dict[str, str]) -> str | None:
    normalized = normalize_name(name)
    return name_index.get(normalized) or NAME_ALIASES.get(normalized)


def fetch_sovereign_scope(name_index: dict[str, str]) -> tuple[set[str], list[str], list[str]]:
    """Return the 193 UN members plus the Holy See and State of Palestine."""
    source = fetch_text(UN_MEMBER_STATES_URL)
    (RAW_DIR / "un-member-states.html").write_text(source, encoding="utf-8")
    tree = html.fromstring(source)
    names = [
        node.text_content().strip()
        for node in tree.xpath("//h2")
        if node.text_content().strip() != "Search the United Nations"
    ]
    member_codes: set[str] = set()
    unmatched: list[str] = []
    for name in names:
        code = resolve_name(name, name_index)
        if code is None:
            unmatched.append(name)
        else:
            member_codes.add(code)

    sovereign_codes = member_codes | {"VAT", "PSE"}
    if len(names) != 193 or len(member_codes) != 193 or len(sovereign_codes) != 195:
        raise RuntimeError(
            "Unexpected UN scope: "
            f"page_names={len(names)}, matched_members={len(member_codes)}, sovereign_total={len(sovereign_codes)}, "
            f"unmatched={unmatched}"
        )
    return sovereign_codes, sorted(unmatched), names


def fetch_world_bank_series(master_codes: set[str]) -> dict[str, dict[str, Observation]]:
    output: dict[str, dict[str, Observation]] = {key: {} for key in WB_INDICATORS}
    for key, meta in WB_INDICATORS.items():
        url = WORLD_BANK_API.format(indicator=meta["code"], start=meta["start"], end=meta["target"])
        payload = fetch_json(url)
        (RAW_DIR / f"world-bank-{key}.json").write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        history: dict[str, dict[int, float]] = defaultdict(dict)
        for item in payload[1] or []:
            code = item.get("countryiso3code")
            value = item.get("value")
            if code in master_codes and value is not None:
                history[code][int(item["date"])] = float(value)

        for code in master_codes:
            years = history.get(code, {})
            target = meta["target"]
            if target in years:
                display_year = meta.get("display_year", target)
                note = "2023 LPI edition; WDI API survey-year field is 2022" if key == "lpi" else ""
                output[key][code] = Observation(years[target], display_year, "observed", url, note)
                continue
            valid = [year for year in years if year <= target and target - year <= meta["max_lag"]]
            if valid:
                year = max(valid)
                display_year = year
                output[key][code] = Observation(years[year], display_year, "substituted_older_year", url, f"Target {meta.get('display_year', target)}; used latest official observation {year}")
            else:
                output[key][code] = Observation(None, None, "missing", url, f"No official observation within {meta['max_lag']} years of target")
    return output


def parse_wipo(name_index: dict[str, str]) -> tuple[dict[str, Observation], list[str]]:
    workbook = load_workbook(WIPO_FILE, read_only=True, data_only=True)
    sheet = workbook["A30"]
    output: dict[str, Observation] = {}
    unmatched: list[str] = []
    for row in sheet.iter_rows(min_row=16, max_row=194, values_only=True):
        name = row[1]
        if not name or name == "Total":
            continue
        if normalize_name(str(name)) == "others":
            continue
        value = to_number(row[4])
        code = resolve_name(str(name), name_index)
        if code is None:
            if value is not None:
                unmatched.append(str(name))
            continue
        output[code] = Observation(
            value,
            2024,
            "official_source_estimate" if value is not None else "missing",
            WIPO_URL,
            "WIPO A30, by country of origin; workbook note states that 2024 data are WIPO estimates",
        )
    workbook.close()
    return output, sorted(set(unmatched))


def parse_wipo_fact_sheet(name_index: dict[str, str]) -> tuple[dict[str, Observation], list[str], list[dict[str, Any]]]:
    if not WIPO_FACT_SHEET_FILE.exists():
        WIPO_FACT_SHEET_FILE.write_bytes(fetch_bytes(WIPO_FACT_SHEET_URL))

    output: dict[str, Observation] = {}
    unmatched: list[str] = []
    raw_rows: list[dict[str, Any]] = []
    with pdfplumber.open(WIPO_FACT_SHEET_FILE) as document:
        for page in document.pages:
            lines: dict[float, list[dict[str, Any]]] = defaultdict(list)
            for word in page.extract_words():
                lines[round(float(word["top"]), 1)].append(word)
            for words in lines.values():
                name_words = [word for word in words if float(word["x0"]) < 145]
                application_words = [
                    word
                    for word in words
                    if 145 <= float(word["x0"]) < 180 and re.fullmatch(r"[\d,]+", str(word["text"]))
                ]
                rank_words = [
                    word
                    for word in words
                    if 180 <= float(word["x0"]) < 210 and re.fullmatch(r"\d+", str(word["text"]))
                ]
                if not name_words or len(application_words) != 1 or len(rank_words) != 1:
                    continue
                name = " ".join(str(word["text"]) for word in sorted(name_words, key=lambda word: float(word["x0"])))
                name = name.replace("T��rkiye", "Türkiye")
                value = int(str(application_words[0]["text"]).replace(",", ""))
                rank = int(rank_words[0]["text"])
                code = resolve_name(name, name_index)
                raw_rows.append({"name": name, "pct_applications_2024": value, "global_rank": rank, "source_url": WIPO_FACT_SHEET_URL})
                if code is None:
                    unmatched.append(name)
                    continue
                output[code] = Observation(
                    value,
                    2024,
                    "observed",
                    WIPO_FACT_SHEET_URL,
                    "WIPO Statistics Database 2024 fact sheet; source updated 7/2026",
                )
    if len(raw_rows) != 123:
        raise RuntimeError(f"Unexpected WIPO PCT positive-origin row count: {len(raw_rows)}")
    return output, sorted(set(unmatched)), raw_rows


def parse_un_electricity(name_index: dict[str, str]) -> tuple[dict[str, Observation], list[str], list[dict[str, Any]]]:
    if not UN_ENERGY_FILE.exists():
        request = urllib.request.Request(
            UN_ENERGY_PDF_URL,
            headers={"User-Agent": USER_AGENT, "Referer": "https://desapublications.un.org/"},
        )
        with urllib.request.urlopen(request, timeout=90) as response:
            UN_ENERGY_FILE.write_bytes(response.read())

    reader = PdfReader(UN_ENERGY_FILE)
    lines: list[str] = []
    # Printed pages 61-67 contain the country indicator table (PDF pages 65-71).
    for page in reader.pages[64:71]:
        lines.extend((page.extract_text() or "").splitlines())

    value_pattern = r"(?:-|\d[\d,]*(?:\.\d+)?)"
    row_pattern = re.compile(rf"^(?P<name>.+?)\s+(?P<values>({value_pattern}\s+){{7}}{value_pattern})$")
    output: dict[str, Observation] = {}
    unmatched: list[str] = []
    raw_rows: list[dict[str, Any]] = []
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if index + 1 < len(lines) and not row_pattern.match(line) and re.match(r"^[\d,-]", lines[index + 1].strip()):
            line += " " + lines[index + 1].strip()
            index += 1
        match = row_pattern.match(line)
        if match:
            name = re.sub(r"(?:14|15|16)$", "", match.group("name")).strip()
            values = match.group("values").split()
            electricity = to_number(values[5])
            code = resolve_name(name, name_index)
            raw_rows.append(
                {
                    "name": name,
                    "electricity_consumption_kwh_per_capita_2022": electricity,
                    "source_url": UN_ENERGY_PDF_URL,
                }
            )
            if code is None:
                unmatched.append(name)
            elif electricity is not None:
                output[code] = Observation(
                    electricity,
                    2022,
                    "substituted_alternative_source",
                    UN_ENERGY_PDF_URL,
                    "UNSD electricity consumption per capita; used only when the WDI electricity series is unavailable",
                )
        index += 1
    if len(raw_rows) != 227:
        raise RuntimeError(f"Unexpected UN energy country-row count: {len(raw_rows)}")
    return output, sorted(set(unmatched)), raw_rows


def compare_electricity_sources(un_electricity: dict[str, Observation]) -> dict[str, Any]:
    payload = json.loads((RAW_DIR / "world-bank-electricity.json").read_text(encoding="utf-8"))
    wb_2022 = {
        item["countryiso3code"]: float(item["value"])
        for item in payload[1] or []
        if item.get("value") is not None and item.get("date") == "2022"
    }
    differences = sorted(
        abs(float(un_electricity[code].value) - value) / value
        for code, value in wb_2022.items()
        if code in un_electricity and value > 0 and un_electricity[code].value is not None
    )
    if not differences:
        return {"overlap": 0, "median_absolute_percent_difference": None, "p90_absolute_percent_difference": None}
    median = differences[len(differences) // 2] if len(differences) % 2 else (differences[len(differences) // 2 - 1] + differences[len(differences) // 2]) / 2
    return {
        "overlap": len(differences),
        "median_absolute_percent_difference": median * 100,
        "p90_absolute_percent_difference": differences[int(0.9 * len(differences))] * 100,
    }


def parse_nature(name_index: dict[str, str]) -> tuple[dict[str, Observation], list[str], list[dict[str, Any]]]:
    output: dict[str, Observation] = {}
    unmatched: list[str] = []
    raw_rows: list[dict[str, Any]] = []
    for scope in NATURE_SCOPES:
        url = NATURE_BASE.format(scope=scope)
        source = fetch_text(url)
        tables = pd.read_html(io.StringIO(source))
        if len(tables) != 1:
            raise RuntimeError(f"Unexpected Nature table count for {scope}: {len(tables)}")
        table = tables[0]
        for record in table.to_dict("records"):
            name = str(record["Country/territory"])
            value = to_number(record["Share 2024"])
            code = resolve_name(name, name_index)
            raw_rows.append({"scope": urllib.parse.unquote(scope), "name": name, "share_2024": value, "source_url": url})
            if code is None:
                unmatched.append(name)
                continue
            if value is None:
                continue
            existing = output.get(code)
            if existing and not math.isclose(float(existing.value), float(value), rel_tol=0, abs_tol=1e-9):
                raise RuntimeError(f"Conflicting Nature values for {code}: {existing.value} vs {value}")
            output[code] = Observation(value, 2024, "observed", url, "Nature Index 2025 Research Leaders regional table")
    return output, sorted(set(unmatched)), raw_rows


def clarify_hcr_page(region: str, page: int) -> str:
    params = urllib.parse.urlencode(
        {
            "action": "clv_hcr_members_filter",
            "clv-paged": page,
            "clv-region": region,
            "filter": "hcr-member",
        }
    )
    url = f"{CLARIVATE_AJAX_URL}?{params}"
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT, "Referer": CLARIVATE_LIST_URL})
    last_error: Exception | None = None
    for attempt in range(3):
        try:
            with urllib.request.urlopen(request, timeout=60) as response:
                return response.read().decode("utf-8", errors="replace")
        except Exception as error:  # noqa: BLE001
            last_error = error
            time.sleep(1.5 * (attempt + 1))
    raise RuntimeError(f"Failed Clarivate query for {region}, page {page}") from last_error


def hcr_page_summary(region: str) -> tuple[str, int, int]:
    first = clarify_hcr_page(region, 1)
    tree = html.fromstring(first)
    rows = len(tree.xpath("//tbody[contains(concat(' ', normalize-space(@class), ' '), ' hcr-member ')]"))
    pages = [int(value) for value in tree.xpath("//*[@data-page]/@data-page") if str(value).isdigit()]
    last_page = max(pages, default=1)
    if last_page == 1:
        return region, rows, 1
    last = clarify_hcr_page(region, last_page)
    last_tree = html.fromstring(last)
    last_rows = len(last_tree.xpath("//tbody[contains(concat(' ', normalize-space(@class), ' '), ' hcr-member ')]"))
    return region, (last_page - 1) * 10 + last_rows, last_page


def parse_clarivate(name_index: dict[str, str]) -> tuple[dict[str, Observation], list[str], list[dict[str, Any]]]:
    cache_path = RAW_DIR / "clarivate-hcr-2025-counts.csv"
    if cache_path.exists():
        with cache_path.open("r", encoding="utf-8-sig", newline="") as handle:
            cached = list(csv.DictReader(handle))
        if cached:
            output: dict[str, Observation] = {}
            unmatched: list[str] = []
            for row in cached:
                region = row["region"]
                code = resolve_name(region, name_index)
                if code is None:
                    unmatched.append(region)
                    continue
                output[code] = Observation(int(row["current_list_entries"]), 2025, "observed", CLARIVATE_LIST_URL, "Current online list entries; cached from this audit date")
            return output, sorted(set(unmatched)), cached

    main_source = fetch_text(CLARIVATE_LIST_URL)
    tree = html.fromstring(main_source)
    regions = [
        value.strip()
        for value in tree.xpath("//select[@name='clv-region']/option/@value")
        if value.strip() and value.strip() != "N/A"
    ]
    results: list[tuple[str, int, int]] = []
    with ThreadPoolExecutor(max_workers=6) as executor:
        futures = {executor.submit(hcr_page_summary, region): region for region in regions}
        for future in as_completed(futures):
            results.append(future.result())
    results.sort(key=lambda item: item[0])

    output: dict[str, Observation] = {}
    unmatched: list[str] = []
    raw_rows: list[dict[str, Any]] = []
    for region, count, pages in results:
        code = resolve_name(region, name_index)
        raw_rows.append({"region": region, "current_list_entries": count, "pages": pages, "source_url": CLARIVATE_LIST_URL})
        if code is None:
            unmatched.append(region)
            continue
        output[code] = Observation(count, 2025, "observed", CLARIVATE_LIST_URL, "Current online list entries; live list may differ from annual press-release totals")
    return output, sorted(set(unmatched)), raw_rows


def clip(value: float) -> float:
    return min(100.0, max(0.0, value))


def log_goalpost(value: float, lower: float, upper: float) -> float:
    return clip(100 * math.log(value / lower) / math.log(upper / lower))


def log_count(value: float, upper: float) -> float:
    return clip(100 * math.log1p(value) / math.log1p(upper))


def linear(value: float, lower: float, upper: float) -> float:
    return clip(100 * (value - lower) / (upper - lower))


def calculate_score(values: dict[str, float]) -> tuple[float, dict[str, float]]:
    economy = log_goalpost(values["gni"], 1000, 80000)
    knowledge = (
        0.4 * log_count(values["pct"], 100000)
        + 0.4 * log_count(values["nature"], 50000)
        + 0.2 * log_count(values["hcr"], 3000)
    )
    health = linear(values["life"], 20, 85)
    infrastructure = (
        log_goalpost(values["electricity"], 500, 10000)
        + clip(values["internet"])
        + linear(values["lpi"], 1, 5)
        + clip(values["water"])
    ) / 4
    total = 0.3 * economy + 0.3 * knowledge + 0.2 * health + 0.2 * infrastructure
    return total, {"economy": economy, "knowledge": knowledge, "health": health, "infrastructure": infrastructure}


def validate_observation(key: str, observation: Observation) -> str | None:
    value = observation.value
    if value is None:
        return None
    rules = {
        "gni": (0, None),
        "pct": (0, None),
        "nature": (0, None),
        "hcr": (0, None),
        "life": (20, 100),
        "electricity": (0, None),
        "internet": (0, 100),
        "lpi": (1, 5),
        "water": (0, 100),
    }
    lower, upper = rules[key]
    if float(value) < lower or (upper is not None and float(value) > upper):
        return f"{key}={value} outside [{lower}, {upper}]"
    return None


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    columns = list(rows[0].keys())
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    ensure_dirs()
    master = build_master()
    master_codes = {row["iso3"] for row in master}
    name_index = build_name_index(master)
    sovereign_codes, un_unmatched, un_member_names = fetch_sovereign_scope(name_index)
    for row in master:
        row["scope"] = "sovereign_195" if row["iso3"] in sovereign_codes else "extended_economy_or_territory"
    (RAW_DIR / "un-member-state-names.json").write_text(
        json.dumps(
            {
                "source": UN_MEMBER_STATES_URL,
                "members": un_member_names,
                "non_member_observers_added": ["Holy See", "State of Palestine"],
                "observer_source": UN_OBSERVER_STATES_URL,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    wb = fetch_world_bank_series(master_codes)
    pct, pct_unmatched, pct_raw = parse_wipo_fact_sheet(name_index)
    un_electricity, un_electricity_unmatched, un_electricity_raw = parse_un_electricity(name_index)
    electricity_comparison = compare_electricity_sources(un_electricity)
    for code, replacement in un_electricity.items():
        if code in wb["electricity"] and wb["electricity"][code].value is None:
            wb["electricity"][code] = replacement
    nature, nature_unmatched, nature_raw = parse_nature(name_index)
    hcr, hcr_unmatched, hcr_raw = parse_clarivate(name_index)

    # For complete non-zero source tables, absence means no recorded output/list entry.
    for code in master_codes:
        nature.setdefault(code, Observation(0, 2024, "inferred_zero", NATURE_BASE.format(scope="global"), "No recorded 2024 Share in any regional Nature Index table"))
        hcr.setdefault(code, Observation(0, 2025, "inferred_zero", CLARIVATE_LIST_URL, "Country/region absent from the complete 2025 filter list"))
        pct.setdefault(
            code,
            Observation(
                0,
                2024,
                "inferred_zero",
                WIPO_FACT_SHEET_URL,
                "No positive 2024 PCT applications in WIPO's complete country-of-origin ranking table",
            ),
        )

    all_series: dict[str, dict[str, Observation]] = {
        "gni": wb["gni"],
        "pct": pct,
        "nature": nature,
        "hcr": hcr,
        "life": wb["life"],
        "electricity": wb["electricity"],
        "internet": wb["internet"],
        "lpi": wb["lpi"],
        "water": wb["water"],
    }

    invalid: list[str] = []
    output_rows: list[dict[str, Any]] = []
    output_json: list[dict[str, Any]] = []
    for country in master:
        code = country["iso3"]
        observations = {key: all_series[key][code] for key in all_series}
        for key, observation in observations.items():
            issue = validate_observation(key, observation)
            if issue:
                invalid.append(f"{code}: {issue}")
                observation.value = None
                observation.status = "invalid"
                observation.note = issue

        missing = [key for key, item in observations.items() if item.value is None]
        statuses = {item.status for item in observations.values()}
        if missing:
            score_status = "insufficient_data"
            score = None
            dimensions = None
        else:
            values = {key: float(item.value) for key, item in observations.items()}
            score, dimensions = calculate_score(values)
            if statuses <= {"observed"}:
                score_status = "fully_observed"
            else:
                score_status = "calculated_with_substitution_or_inference"

        flat: dict[str, Any] = {
            "iso3": code,
            "name": country["name"],
            "scope": country["scope"],
            "region": country["region"],
            "income": country["income"],
            "score_status": score_status,
            "cdi": round(score, 4) if score is not None else None,
            "missing_indicators": ";".join(missing),
        }
        for key, item in observations.items():
            flat[f"{key}_value"] = item.value
            flat[f"{key}_year"] = item.year
            flat[f"{key}_status"] = item.status
        output_rows.append(flat)
        output_json.append(
            {
                **country,
                "score_status": score_status,
                "cdi": round(score, 6) if score is not None else None,
                "dimensions": {key: round(value, 6) for key, value in (dimensions or {}).items()} if dimensions else None,
                "missing_indicators": missing,
                "indicators": {
                    key: {
                        "value": item.value,
                        "year": item.year,
                        "status": item.status,
                        "source": item.source,
                        "note": item.note,
                    }
                    for key, item in observations.items()
                },
            }
        )

    write_csv(RAW_DIR / "nature-index-share-2024.csv", nature_raw)
    write_csv(RAW_DIR / "clarivate-hcr-2025-counts.csv", hcr_raw)
    write_csv(RAW_DIR / "un-electricity-consumption-per-capita-2022.csv", un_electricity_raw)
    write_csv(RAW_DIR / "wipo-pct-positive-origins-2024.csv", pct_raw)
    write_csv(PROCESSED_DIR / "global-cdi-coverage.csv", output_rows)
    (PROCESSED_DIR / "global-cdi-coverage.json").write_text(
        json.dumps(
            {
                "snapshot": {
                    "accessed": ACCESS_DATE,
                    "master_count": len(master),
                    "sovereign_country_count": len(sovereign_codes),
                    "formula_version": "CDI 1.0 base score",
                    "policy": {
                        "world_bank_fallback": "GNI/life: two years; internet/water: five years; every fallback is labeled",
                        "lpi_fallback": "latest official 2014-2018 LPI 1.0 value when the 2023 edition is unavailable",
                        "electricity_alternative": "UNSD 2022 electricity consumption per capita only when WDI is unavailable",
                        "pct_2024": "current WIPO Statistics Database fact sheet; absence from complete positive-origin table is inferred zero",
                        "nature_absence": "inferred zero when absent from all complete regional non-zero tables",
                        "hcr_absence": "inferred zero when absent from the complete country/region filter list",
                        "cross_indicator_proxy": "not allowed",
                    },
                },
                "countries": output_json,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    def summarize_scope(scope_codes: set[str]) -> dict[str, Any]:
        scoped_coverage: dict[str, Any] = {}
        for key, series in all_series.items():
            observations = [series[code] for code in scope_codes]
            counts = Counter(item.status for item in observations)
            scoped_coverage[key] = {
                "available": sum(1 for item in observations if item.value is not None),
                "missing": sum(1 for item in observations if item.value is None),
                "statuses": dict(sorted(counts.items())),
            }
        scoped_rows = [row for row in output_rows if row["iso3"] in scope_codes]
        scoped_score_counts = Counter(row["score_status"] for row in scoped_rows)
        scoped_missing = Counter()
        for row in scoped_rows:
            for key in filter(None, row["missing_indicators"].split(";")):
                scoped_missing[key] += 1
        return {
            "count": len(scope_codes),
            "score_status": dict(scoped_score_counts),
            "indicator_coverage": scoped_coverage,
            "missing_by_indicator": dict(scoped_missing),
        }

    sovereign_summary = summarize_scope(sovereign_codes)
    extended_summary = summarize_scope(master_codes)
    coverage = sovereign_summary["indicator_coverage"]
    score_counts = Counter(sovereign_summary["score_status"])
    missing_by_indicator = Counter(sovereign_summary["missing_by_indicator"])

    report = [
        "# CDI 全球数据覆盖审计",
        "",
        f"- 审计日期：{ACCESS_DATE}",
        "- 主口径：195 个国家（193 个联合国会员国 + 圣座 + 巴勒斯坦国）。",
        f"- 扩展口径：{len(master)} 个国家、经济体或属地（世界银行 217 个条目，另加入台湾和梵蒂冈）。",
        "- 基础分公式版本：CDI 1.0；应用层公式版本：CDI 1.1。",
        "- 本报告审计九项基础指标；资源依赖修正数据由 scripts/update_resource_rents.mjs 单独生成。",
        "",
        "## 195 国主口径的计算资格",
        "",
        f"- 全部为直接观测：{score_counts.get('fully_observed', 0)}",
        f"- 含旧年份替代值或推定零、仍可计算：{score_counts.get('calculated_with_substitution_or_inference', 0)}",
        f"- 至少一项必要数据缺失、不计算总分：{score_counts.get('insufficient_data', 0)}",
        "",
        "## 195 国主口径的各指标覆盖",
        "",
        "| 指标 | 可用 | 缺失 | 状态构成 |",
        "|---|---:|---:|---|",
    ]
    for key in all_series:
        item = coverage[key]
        status_text = ", ".join(f"{name}={count}" for name, count in item["statuses"].items())
        report.append(f"| {key} | {item['available']} | {item['missing']} | {status_text} |")

    report.extend(["", "## 导致无法计算的缺失项", ""])
    for key, count in missing_by_indicator.most_common():
        report.append(f"- {key}: {count} 个国家/经济体")
    if not missing_by_indicator:
        report.append("- 无")

    report.extend(
        [
            "",
            "## 使用规则",
            "",
            "1. 世界银行连续指标优先采用目标年份。GNI 与寿命最多回退两年；互联网与基础饮水最多回退五年；全部标记 `substituted_older_year`。",
            "2. LPI 以 2023 版为目标。WDI API 将该版记录在 2022 调查年；缺失时允许采用该国最后一个 2014-2018 年 LPI 1.0 官方值，并明确显示年份。更早数据不采用。",
            "3. 世界银行人均用电缺失时，使用联合国统计司 2022 年人均电力消费作为替代来源，标记 `substituted_alternative_source`。两来源 2022 年重叠样本的中位绝对差见下方质量检查。",
            "4. Nature Index 区域表列到极低的非零 Share。未出现在任何区域表中的经济体记为 `inferred_zero`，不是直接观测。",
            "5. Clarivate 当前名单的国家/地区筛选项覆盖所有有名单记录的地区。筛选项中不存在的经济体记为 `inferred_zero`。",
            "6. PCT 采用 WIPO Statistics Database 的当前 2024 Fact Sheet（来源更新 7/2026），不用旧 Yearly Review 工作簿中的初步估计。正值国家完整排名表之外的国家记为 `inferred_zero`。",
            "7. 不允许用不同含义的指标替代 PCT、Nature、HCR 或 LPI。仍有缺失项时不计算 CDI。",
            "",
            "## 替代来源质量检查",
            "",
            f"- 世界银行与 UNSD 2022 年人均用电重叠样本：{electricity_comparison['overlap']} 个经济体。",
            f"- 中位绝对百分比差：{electricity_comparison['median_absolute_percent_difference']:.1f}%；90 分位绝对差：{electricity_comparison['p90_absolute_percent_difference']:.1f}%。",
            "- 因定义和编制方法并非完全相同，UNSD 值只用于 WDI 缺失项，且不会被标成世界银行观测值。",
            "",
            "## 名称映射检查",
            "",
            f"- 联合国会员国未匹配名称：{', '.join(un_unmatched) if un_unmatched else '无'}",
            f"- WIPO 未匹配且有数值的名称：{', '.join(pct_unmatched) if pct_unmatched else '无'}",
            f"- UNSD 人均用电未匹配名称：{', '.join(un_electricity_unmatched) if un_electricity_unmatched else '无'}",
            f"- Nature Index 未匹配名称：{', '.join(nature_unmatched) if nature_unmatched else '无'}",
            f"- Clarivate 未匹配名称：{', '.join(hcr_unmatched) if hcr_unmatched else '无'}",
            f"- 超出合理范围并转为缺失的值：{'; '.join(invalid) if invalid else '无'}",
            "",
            "## 输出",
            "",
            "- `data/processed/global-cdi-coverage.csv`：覆盖审计宽表。",
            "- `data/processed/global-cdi-coverage.json`：后续网页使用的带状态数据。",
            "- `data/raw/`：下载或解析后的源数据快照。",
        ]
    )
    (AUDIT_DIR / "global-coverage-report.md").write_text("\n".join(report) + "\n", encoding="utf-8")

    summary = {
        "headline_scope": "sovereign_195",
        "sovereign_195": sovereign_summary,
        "extended_master": extended_summary,
        "electricity_source_comparison": electricity_comparison,
        "unmatched": {
            "un_members": un_unmatched,
            "wipo": pct_unmatched,
            "un_electricity": un_electricity_unmatched,
            "nature": nature_unmatched,
            "clarivate": hcr_unmatched,
        },
        "invalid": invalid,
    }
    (AUDIT_DIR / "coverage-summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
